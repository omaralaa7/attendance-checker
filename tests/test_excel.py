import shutil

from openpyxl import load_workbook

from app import excel


def _marks(expected):
    present = set(expected["present_serials"])
    marks = {s: ("p" if s in present else "a") for s in range(1, 99)}
    doubts = {s: "faded" for s in expected["doubtful_serials_marked_absent"]}
    return marks, doubts


def test_parse_roster(roster_path):
    r = excel.load_roster(roster_path)
    assert len(r.students) == 98
    assert r.header_row == 7 and r.name_col == 3 and r.serial_col == 1
    assert [s.serial for s in r.students] == list(range(1, 99))
    assert r.students[0].row == 8
    assert r.students[0].name == "احمد جلال عبدالله عبيد عبدالله"
    assert all("\u00a0" not in s.name and s.name == s.name.strip() for s in r.students)
    assert "القوى" in r.title
    assert r.date_cols == {}


def test_write_column(roster_path, expected, tmp_path):
    p = tmp_path / "w.xlsx"
    shutil.copy(roster_path, p)
    before = load_workbook(roster_path).worksheets[0]
    marks, doubts = _marks(expected)

    col = excel.write_column(p, "2026-10-11", marks, doubts)
    assert col == 5

    ws = load_workbook(p).worksheets[0]
    assert ws.cell(7, 5).value == "11/10/2026"
    vals = [ws.cell(r, 5).value for r in range(8, 106)]
    assert vals.count("✓") == 83 and vals.count("✗") == 15
    assert ws.cell(8, 5).fill.fgColor.rgb.endswith("C6EFCE")
    assert ws.cell(10, 5).fill.fgColor.rgb.endswith("FFC7CE")
    assert ws.cell(30 + 7, 5).fill.fgColor.rgb.endswith("FFEB9C")
    assert ws.cell(40 + 7, 5).value == "✗"
    assert ws.cell(7, 5).font.b and ws.cell(7, 5).fill.fgColor.rgb == ws.cell(7, 4).fill.fgColor.rgb
    assert ws.column_dimensions["E"].width == 14

    # title block, merges and original cells untouched
    assert {str(m) for m in ws.merged_cells.ranges} == {str(m) for m in before.merged_cells.ranges}
    for r in range(1, 106):
        for c in range(1, 5):
            assert ws.cell(r, c).value == before.cell(r, c).value

    r = excel.load_roster(p)
    assert r.date_cols == {"2026-10-11": 5}
    got = excel.read_marks(p, "2026-10-11")
    assert sum(v == "p" for v in got["marks"].values()) == 83
    assert set(got["doubts"]) == {30, 40}


def test_same_date_overwrites(roster_path, tmp_path):
    p = tmp_path / "w.xlsx"
    shutil.copy(roster_path, p)
    excel.write_column(p, "2026-10-11", {s: "a" for s in range(1, 99)})
    excel.write_column(p, "2026-10-18", {s: "p" for s in range(1, 99)})
    col = excel.write_column(p, "2026-10-11", {s: "p" for s in range(1, 99)})
    assert col == 5
    ws = load_workbook(p).worksheets[0]
    assert ws.max_column == 6
    assert ws.cell(7, 6).value == "18/10/2026"
    assert all(ws.cell(r, 5).value == "✓" for r in range(8, 106))


def test_matches_expected_output(roster_path, expected, tmp_path):
    ref = roster_path.parent / "expected_output_2026-10-11.xlsx"
    p = tmp_path / "w.xlsx"
    shutil.copy(roster_path, p)
    excel.write_column(p, "2026-10-11", *_marks(expected))
    a, b = load_workbook(p).worksheets[0], load_workbook(ref).worksheets[0]
    for r in range(7, 106):
        assert a.cell(r, 5).value == b.cell(r, 5).value, r
        assert a.cell(r, 5).fill.fgColor.rgb[-6:] == b.cell(r, 5).fill.fgColor.rgb[-6:], r


def test_bad_file(tmp_path):
    import pytest
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"not excel")
    with pytest.raises(excel.RosterError):
        excel.load_roster(p)
