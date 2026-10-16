"""Roster parsing and writing the attendance date column (openpyxl)."""
from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRESENT, ABSENT = "✓", "✗"
GREEN_FONT, GREEN_FILL = "FF006100", "FFC6EFCE"
RED_FONT, RED_FILL = "FF9C0006", "FFFFC7CE"
YELLOW_FILL = "FFFFEB9C"

_DATE_RE = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})$")


class RosterError(ValueError):
    """Raised with an Arabic, user-facing message when the workbook is unusable."""


def clean(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%d/%m/%Y")
    return str(v).replace(" ", " ").strip()


def fmt_date(iso: str) -> str:
    """'2026-10-11' -> '11/10/2026'"""
    y, m, d = iso.split("-")
    return f"{int(d)}/{int(m)}/{y}"


def header_to_iso(v) -> str | None:
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    m = _DATE_RE.match(clean(v))
    if not m:
        return None
    d, mo, y = int(m[1]), int(m[2]), int(m[3])
    try:
        return date(y, mo, d).isoformat()
    except ValueError:
        return None


@dataclass
class Student:
    serial: int
    row: int
    name: str


@dataclass
class Roster:
    title: str
    header_row: int
    name_col: int
    serial_col: int
    students: list[Student]
    date_cols: dict[str, int] = field(default_factory=dict)

    def to_json(self) -> dict:
        return {
            "title": self.title,
            "students": [s.__dict__ for s in self.students],
            "dates": sorted(self.date_cols),
        }


def parse_ws(ws) -> Roster:
    header_row = name_col = 0
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 50)):
        for c in row:
            if "اسم الطالب" in clean(c.value):
                header_row, name_col = c.row, c.column
                break
        if header_row:
            break
    if not header_row:
        raise RosterError('مش لاقي عمود "اسم الطالب" في الملف.')

    serial_col = 1
    date_cols: dict[str, int] = {}
    for c in ws[header_row]:
        t = clean(c.value)
        if t == "م":
            serial_col = c.column
        iso = header_to_iso(c.value)
        if iso:
            date_cols[iso] = c.column

    students: list[Student] = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = " ".join(clean(ws.cell(r, name_col).value).split())
        if not name:
            continue
        s = clean(ws.cell(r, serial_col).value)
        serial = int(s) if s.isdigit() else len(students) + 1
        students.append(Student(serial=serial, row=r, name=name))
    if not students:
        raise RosterError("الملف مفيهوش أسماء طلبة تحت عمود \"اسم الطالب\".")

    return Roster(
        title=clean(ws["B2"].value),
        header_row=header_row,
        name_col=name_col,
        serial_col=serial_col,
        students=students,
        date_cols=date_cols,
    )


def load_roster(path: str | Path) -> Roster:
    try:
        wb = load_workbook(path)
    except Exception as e:  # noqa: BLE001 - any parse error is "not a readable xlsx"
        raise RosterError("الملف ده مش ملف إكسل مقروء (.xlsx).") from e
    return parse_ws(wb.worksheets[0])


def read_marks(path: str | Path, iso: str) -> dict:
    """Marks already stored for a date: {marks:{serial:'p'|'a'}, doubts:{serial:''}}."""
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    roster = parse_ws(ws)
    col = roster.date_cols.get(iso)
    marks: dict[int, str] = {}
    doubts: dict[int, str] = {}
    if col:
        for s in roster.students:
            c = ws.cell(s.row, col)
            v = clean(c.value)
            if v == PRESENT:
                marks[s.serial] = "p"
            elif v == ABSENT:
                marks[s.serial] = "a"
            rgb = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == "solid" else None
            if isinstance(rgb, str) and rgb.upper().endswith("FFEB9C"):
                doubts[s.serial] = ""
    return {"marks": marks, "doubts": doubts, "exists": bool(col)}


def write_column(path: str | Path, iso: str, marks: dict, doubts: dict | None = None,
                 out_path: str | Path | None = None) -> int:
    """Write (or overwrite) the attendance column for `iso`. Keys of marks/doubts are serials.

    Returns the column index written.
    """
    doubts = {int(k): v for k, v in (doubts or {}).items()}
    marks = {int(k): v for k, v in marks.items()}
    wb = load_workbook(path)
    ws = wb.worksheets[0]
    roster = parse_ws(ws)

    col = roster.date_cols.get(iso)
    if not col:
        last = max((c.column for c in ws[roster.header_row] if clean(c.value)), default=0)
        col = last + 1

    hdr = ws.cell(roster.header_row, col)
    ref = ws.cell(roster.header_row, max(1, col - 1))
    hdr.value = fmt_date(iso)
    if ref.has_style and ref.column != col:
        hdr._style = copy(ref._style)
    hdr.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(col)].width = 14

    for s in roster.students:
        c = ws.cell(s.row, col)
        m = marks.get(s.serial)
        name_cell = ws.cell(s.row, roster.name_col)
        c.value = PRESENT if m == "p" else ABSENT if m == "a" else None
        c.border = copy(name_cell.border)
        c.alignment = Alignment(horizontal="center", vertical="center")
        if c.value is None:
            c.font = Font()
            c.fill = PatternFill(fill_type=None)
            continue
        present = m == "p"
        c.font = Font(bold=True, size=14, color=GREEN_FONT if present else RED_FONT)
        fill = YELLOW_FILL if s.serial in doubts else GREEN_FILL if present else RED_FILL
        c.fill = PatternFill(fill_type="solid", start_color=fill, end_color=fill)

    wb.save(out_path or path)
    return col
